"""Chat file attachments → text a model can read (1.2-M5, D41).

Everything here is local and passive: a file is sniffed, decoded and its
text pulled out by pure-Python readers — nothing in it is executed (no
macros, no formulas re-evaluated, no external links followed: openpyxl runs
with ``keep_links=False`` and ``data_only=True``, so a formula cell gives the
value Excel last saved). Paths come from the upload / works index, never
from a request; the file name is only used for display and as a type hint.

What a file becomes is an ``Extracted``: its type, whether there is text to
read (``readable``, with a ``reason`` when not), the text itself, where each
page / slide / sheet sits in it (``segments``) and ``info`` — the facts the
page shows next to the attachment (``pages``, ``sheets``, ``rows``,
``slides``, ``chars``, ``encoding``…, and ``truncated`` with a note whenever
less than the whole file was taken). Results are cached by (path, mtime,
size): the same file is never extracted twice.

Guards (決策記錄 1.2-M5-c): a text file is read up to ``MAX_TEXT_READ_BYTES``;
a zip-based Office file is checked for entry count, total uncompressed size
and compression ratio before any library opens it (a decompression bomb
never gets unpacked); spreadsheets stop at ``MAX_SHEET_ROWS`` rows /
``MAX_SHEET_COLS`` columns per sheet; every extraction stops at
``MAX_EXTRACT_CHARS`` characters and runs under ``EXTRACT_TIMEOUT_S``.

``read_range`` / ``search`` work on an ``Extracted`` and back the chat's
``read_file`` / ``search_file`` tools.
"""

from __future__ import annotations

import asyncio
import codecs
import csv
import io
import logging
import re
import threading
import zipfile
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from .. import modalities as _MOD

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------- limits (決策記錄 1.2-M5-c)
#: bytes of a text file that are read (the rest is reported as truncated)
MAX_TEXT_READ_BYTES = 8 * 1024 * 1024
#: characters kept from any one file
MAX_EXTRACT_CHARS = 2_000_000
#: PDF pages whose text is extracted
MAX_PDF_PAGES = 2000
#: a zip-based Office file (docx / xlsx / pptx) is refused past any of these
MAX_ZIP_ENTRIES = 10_000
MAX_ZIP_UNCOMPRESSED = 300 * 1024 * 1024
MAX_ZIP_RATIO = 200  # per member, for members over 1 MiB uncompressed
#: spreadsheet limits per sheet, and sheets per workbook
MAX_SHEET_ROWS = 20_000
MAX_SHEET_COLS = 100
MAX_SHEETS = 50
#: slides read from one deck
MAX_SLIDES = 1000
#: one extraction gives up after this long (the thread finishes and caches anyway)
EXTRACT_TIMEOUT_S = 60.0
#: bytes looked at to tell what a file is
SNIFF_BYTES = 8192
_CACHE_CHARS = 40_000_000  # extracted text kept in memory

TEXT_EXTS = {
    ".txt": "text", ".text": "text", ".log": "text", ".md": "markdown", ".markdown": "markdown", ".rst": "text",
    ".csv": "csv", ".tsv": "csv", ".json": "json", ".jsonl": "json", ".ndjson": "json", ".html": "html", ".htm": "html",
    ".xml": "xml", ".svg": "xml", ".yaml": "text", ".yml": "text", ".toml": "text", ".ini": "text", ".cfg": "text",
    ".conf": "text", ".env": "text", ".srt": "text", ".vtt": "text", ".tex": "text", ".sql": "code",
}
CODE_EXTS = {
    ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".java", ".kt", ".kts", ".c", ".h", ".cc", ".cpp", ".hpp", ".cs",
    ".go", ".rs", ".rb", ".php", ".swift", ".m", ".mm", ".scala", ".lua", ".pl", ".r", ".sh", ".bash", ".zsh", ".ps1",
    ".bat", ".cmd", ".vue", ".svelte", ".css", ".scss", ".less", ".dart", ".ex", ".exs", ".hs", ".clj", ".gradle",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic", ".avif"}
#: an attached .mp4 is offered for transcription like any audio (its audio track is what is read)
AUDIO_EXTS = _MOD.CHAT_AUDIO_EXTS
LEGACY_OFFICE_EXTS = {".doc", ".xls", ".ppt"}

#: type → the word the model and the page read
TYPE_NAMES = {
    "text": "文字", "markdown": "Markdown", "csv": "CSV", "json": "JSON", "html": "HTML", "xml": "XML", "code": "程式碼",
    "pdf": "PDF", "docx": "Word", "xlsx": "Excel", "pptx": "PowerPoint", "audio": "音檔", "image": "圖片",
    "binary": "二進位檔", "legacy_office": "舊版 Office 檔", "transcript": "逐字稿", "lyrics": "歌詞",
}
#: why a file has no text → what the model and the owner are told
REASONS = {
    "binary": "未知的二進位格式，沒有辦法轉成文字",
    "encrypted": "檔案有密碼保護，打不開",
    "corrupt": "檔案損壞或格式不對，打不開",
    "no_text": "沒有可抽取的文字（可能是掃描檔或只有圖片）",
    "empty": "檔案是空的",
    "too_large": "檔案太大或結構太複雜，為了安全沒有展開",
    "timeout": "抽取花太久，已停止",
    "unsupported": "舊版 Office 格式（.doc／.xls／.ppt）讀不了，請另存成新版格式（.docx／.xlsx／.pptx）",
    "missing": "檔案已不在",
    "audio": "音檔沒有文字，要先轉錄",
    "image": "圖片沒有文字",
}
TEXT_TYPES = frozenset({"text", "markdown", "csv", "json", "html", "xml", "code"})


@dataclass
class Segment:
    """Where one page / slide / sheet sits in ``Extracted.text``."""

    unit: str  # page | slide | sheet
    index: int  # 1-based
    start: int
    end: int
    name: str = ""  # a sheet's name
    rows: Optional[int] = None  # a sheet's rows in the text


@dataclass
class Extracted:
    type: str
    readable: bool
    text: str = ""
    reason: Optional[str] = None
    segments: list[Segment] = field(default_factory=list)
    info: dict[str, Any] = field(default_factory=dict)
    _line_starts: Optional[list[int]] = field(default=None, repr=False)

    def __post_init__(self) -> None:
        self.info.setdefault("type", self.type)
        self.info.setdefault("type_name", TYPE_NAMES.get(self.type, self.type))
        self.info["readable"] = self.readable
        self.info["chars"] = len(self.text)
        if self.reason:
            self.info["reason"] = self.reason
            self.info.setdefault("reason_text", REASONS.get(self.reason, self.reason))

    @property
    def chars(self) -> int:
        return len(self.text)

    def line_starts(self) -> list[int]:
        if self._line_starts is None:
            self._line_starts = [0] + [m.end() for m in re.finditer("\n", self.text)]
        return self._line_starts


def unreadable(type_: str, reason: str, **info: Any) -> Extracted:
    return Extracted(type=type_, readable=False, reason=reason, info=dict(info))


# ---------------------------------------------------------------- what a file is
def _zip_kind(path: Path) -> Optional[str]:
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
    except (zipfile.BadZipFile, OSError, ValueError):
        return None
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return "zip"


def _audio_magic(head: bytes) -> bool:
    return (head[:4] == b"RIFF" and head[8:12] == b"WAVE") or head[:3] == b"ID3" or head[:4] in (b"fLaC", b"OggS") \
        or (len(head) > 1 and head[0] == 0xFF and head[1] & 0xE0 == 0xE0) or head[4:8] == b"ftyp" and head[8:11] in (b"M4A", b"mp4", b"iso")


def _image_magic(head: bytes) -> bool:
    return head[:8] == b"\x89PNG\r\n\x1a\n" or head[:3] == b"\xff\xd8\xff" or head[:6] in (b"GIF87a", b"GIF89a") \
        or (head[:4] == b"RIFF" and head[8:12] == b"WEBP")


def audio_format(path: Path) -> Optional[str]:
    """``wav`` / ``mp3`` (the two ``input_audio`` formats every vendor that
    takes audio documents), else ``None``."""
    try:
        with open(path, "rb") as f:
            head = f.read(16)
    except OSError:
        return None
    if head[:4] == b"RIFF" and head[8:12] == b"WAVE":
        return "wav"
    if head[:3] == b"ID3" or (len(head) > 1 and head[0] == 0xFF and head[1] & 0xE0 == 0xE0):
        return "mp3"
    return None


_MP3_KBPS = {  # MPEG-1 Layer III bitrate index → kbps
    1: 32, 2: 40, 3: 48, 4: 56, 5: 64, 6: 80, 7: 96, 8: 112, 9: 128, 10: 160, 11: 192, 12: 224, 13: 256, 14: 320}
#: MPEG-2 / 2.5 Layer III (24 kHz and lower: OpenAI's TTS writes these) bitrate index → kbps
_MP3_LSF_KBPS = {1: 8, 2: 16, 3: 24, 4: 32, 5: 40, 6: 48, 7: 56, 8: 64, 9: 80, 10: 96, 11: 112, 12: 128, 13: 144, 14: 160}
#: second header byte (without the protection bit) → that version's table: MPEG-1, MPEG-2, MPEG-2.5 Layer III
_MP3_TABLES = {0xFA: _MP3_KBPS, 0xF2: _MP3_LSF_KBPS, 0xE2: _MP3_LSF_KBPS}


def audio_duration(path: Path) -> tuple[Optional[float], bool]:
    """``(seconds, estimated)``: exact for WAV, from the first frame's bitrate
    for an MP3 (exact for constant bitrate), ``None`` otherwise — the page
    knows a duration from the browser when this cannot tell."""
    import wave

    fmt = audio_format(path)
    try:
        if fmt == "wav":
            with wave.open(str(path), "rb") as w:
                return round(w.getnframes() / float(w.getframerate() or 1), 2), False
        if fmt == "mp3":
            with open(path, "rb") as f:
                head = f.read(65536)
            i = 0
            if head[:3] == b"ID3" and len(head) > 10:  # skip the tag: its size is 4 syncsafe bytes
                i = 10 + ((head[6] & 0x7F) << 21 | (head[7] & 0x7F) << 14 | (head[8] & 0x7F) << 7 | (head[9] & 0x7F))
                with open(path, "rb") as f:
                    f.seek(i)
                    head, i = f.read(4096), 0
            while i + 4 <= len(head):
                table = _MP3_TABLES.get(head[i + 1] & 0xFE) if head[i] == 0xFF else None  # a Layer III frame
                if table:
                    kbps = table.get(head[i + 2] >> 4)
                    if kbps:
                        return round(path.stat().st_size * 8 / (kbps * 1000), 1), True
                i += 1
    except (OSError, EOFError, ValueError) as e:
        logger.debug("audio duration of %s: %s", path.name, e)
    return None, False


def work_seconds(row: dict[str, Any]) -> Optional[float]:
    """An audio work's length: what the index stored, else read from the file.
    The index only stores a length the tool reported (or a WAV's); a speech
    or music MP3 from the works wall has none, and without it a transcription
    asked about in a chat had no estimate (v1.2.0 acceptance)."""
    if row.get("duration_s"):
        return row["duration_s"]
    p = Path(row["file_path"]) if row.get("file_path") else None
    if p is None or not p.is_file():
        return row.get("duration_s")
    return audio_duration(p)[0]


def audio_info(path: Path) -> dict[str, Any]:
    """What the page shows for an audio upload."""
    info: dict[str, Any] = {"type": "audio", "type_name": TYPE_NAMES["audio"], "format": audio_format(path)}
    seconds, estimated = audio_duration(path)
    if seconds is not None:
        info["duration_s"] = seconds
        if estimated:
            info["duration_estimated"] = True
    return info


def detect(path: Path, name: Optional[str] = None) -> str:
    """The file's type, from its first bytes first and its extension second
    (a renamed file is still what its content says)."""
    ext = Path(name or path.name).suffix.lower()
    with open(path, "rb") as f:
        head = f.read(SNIFF_BYTES)
    if not head:
        return TEXT_EXTS.get(ext) or ("code" if ext in CODE_EXTS else "text")
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head[:4] == b"PK\x03\x04":
        kind = _zip_kind(path)
        return kind if kind in ("docx", "xlsx", "pptx") else "binary"
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        # an OLE container: legacy Office, or a password-protected docx / xlsx / pptx
        return "encrypted_office" if ext in (".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm") else "legacy_office"
    if any(head.startswith(bom) for bom, _ in _BOMS):  # before the audio sniff: FF FE also looks like an MP3 frame sync
        return "code" if ext in CODE_EXTS else TEXT_EXTS.get(ext, "text")
    if _image_magic(head):
        return "image"
    if _audio_magic(head):
        return "audio"
    if _looks_like_text(head):
        if ext in CODE_EXTS:
            return "code"
        return TEXT_EXTS.get(ext, "text")
    if ext in IMAGE_EXTS:
        return "image"
    if ext in AUDIO_EXTS:
        return "audio"
    return "binary"


def _utf16_guess(head: bytes) -> Optional[str]:
    """UTF-16 without a BOM: one byte of every pair is NUL in mostly-ASCII text."""
    sample = head[:2048]
    if len(sample) < 4:
        return None
    even = sample[0::2].count(0) / max(1, len(sample[0::2]))
    odd = sample[1::2].count(0) / max(1, len(sample[1::2]))
    if odd > 0.4 and even < 0.05:
        return "utf-16-le"
    if even > 0.4 and odd < 0.05:
        return "utf-16-be"
    return None


_BOMS = ((codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"), (codecs.BOM_UTF8, "utf-8-sig"),
         (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16"))


def _looks_like_text(head: bytes) -> bool:
    if any(head.startswith(bom) for bom, _ in _BOMS) or _utf16_guess(head):
        return True
    if b"\x00" in head:
        return False
    controls = sum(1 for b in head if b < 32 and b not in (9, 10, 12, 13, 27, 8))
    if controls > len(head) * 0.02:
        return False
    for enc in ("utf-8", "cp950", "gb18030"):
        try:
            codecs.getincrementaldecoder(enc)().decode(head, final=False)
            return True
        except UnicodeDecodeError:
            continue
    # UTF-8 with a few broken bytes is still text (decoded with replacements, marked lossy)
    if head.decode("utf-8", errors="replace").count("�") <= max(2, len(head) // 200):
        return True
    # mostly printable single-byte text (latin-1 family) still counts
    return sum(1 for b in head if b >= 0x80) < len(head) * 0.3


def decode_text(raw: bytes, *, complete: bool = True) -> tuple[str, str, bool]:
    """``(text, encoding, lossy)``. BOMs first, then UTF-16 by its NULs, then
    strict UTF-8, Big5 (cp950), GB18030; anything else is decoded with
    replacement characters and marked ``lossy``. ``complete=False`` (the file
    was cut at ``MAX_TEXT_READ_BYTES``) tolerates a character split at the end."""
    for bom, enc in _BOMS:
        if raw.startswith(bom):
            return raw.decode(enc, errors="replace"), enc.replace("-sig", "") + " (BOM)", False
    guess = _utf16_guess(raw)
    if guess:
        return raw.decode(guess, errors="replace"), guess, False
    for enc in ("utf-8", "cp950", "gb18030"):
        try:
            dec = codecs.getincrementaldecoder(enc)()
            text = dec.decode(raw, final=complete)
            return text, {"cp950": "big5"}.get(enc, enc), False
        except UnicodeDecodeError:
            if enc == "utf-8":
                # UTF-8 with a few broken bytes: keep it UTF-8 (Big5 would turn it all into nonsense)
                loose = raw.decode("utf-8", errors="replace")
                if loose.count("�") <= max(2, len(loose) // 200):
                    return loose, "utf-8", True
            continue
    return raw.decode("utf-8", errors="replace"), "unknown", True


# ---------------------------------------------------------------- extractors
class _Builder:
    """Collects text with page / slide / sheet segments, up to ``MAX_EXTRACT_CHARS``."""

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.size = 0
        self.segments: list[Segment] = []
        self.full = False

    def add(self, s: str) -> None:
        if self.full or not s:
            return
        room = MAX_EXTRACT_CHARS - self.size
        if len(s) > room:
            s, self.full = s[:room], True
        self.parts.append(s)
        self.size += len(s)

    def segment(self, unit: str, index: int, header: str, body: str, **kw: Any) -> None:
        if self.full:
            return
        self.add(header)
        start = self.size
        self.add(body)
        self.segments.append(Segment(unit=unit, index=index, start=start, end=self.size, **kw))
        self.add("\n\n")

    @property
    def text(self) -> str:
        return "".join(self.parts).rstrip("\n")


def _truncated(info: dict[str, Any], note: str) -> None:
    info["truncated"] = True
    info.setdefault("truncated_notes", []).append(note)


def _extract_text(path: Path, type_: str) -> Extracted:
    size = path.stat().st_size
    with open(path, "rb") as f:
        raw = f.read(MAX_TEXT_READ_BYTES)
    if not raw:
        return unreadable(type_, "empty")
    complete = size <= MAX_TEXT_READ_BYTES
    text, encoding, lossy = decode_text(raw, complete=complete)
    info: dict[str, Any] = {"encoding": encoding}
    if lossy:
        info["lossy"] = True  # some bytes became replacement characters
    if not complete:
        _truncated(info, f"檔案 {size:,} bytes，只讀了前 {MAX_TEXT_READ_BYTES:,} bytes")
    if len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS]
        _truncated(info, f"只取了前 {MAX_EXTRACT_CHARS:,} 字")
    text = text.replace("\r\n", "\n")
    if not text.strip():
        return unreadable(type_, "empty", **info)
    lines = text.count("\n") + 1
    info["lines"] = lines
    if type_ == "csv":
        info["rows"] = lines - (1 if text.endswith("\n") else 0)
        try:
            first = next(csv.reader(io.StringIO(text[:65536]), delimiter="\t" if "\t" in text.split("\n", 1)[0] else ","))
            info["cols"] = len(first)
        except (StopIteration, csv.Error):
            pass
    return Extracted(type=type_, readable=True, text=text, info=info)


def _extract_pdf(path: Path) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError, PdfReadError

    info: dict[str, Any] = {}
    with open(path, "rb") as f:
        try:
            reader = PdfReader(f)
            if reader.is_encrypted:
                try:
                    if not reader.decrypt(""):  # owner-password-only PDFs open with an empty user password
                        return unreadable("pdf", "encrypted")
                except Exception:
                    return unreadable("pdf", "encrypted")
            n = len(reader.pages)
        except FileNotDecryptedError:
            return unreadable("pdf", "encrypted")
        except (PdfReadError, ValueError, KeyError, TypeError, OSError, RecursionError) as e:
            logger.info("pdf %s could not be opened: %s", path.name, e)
            return unreadable("pdf", "corrupt")
        info["pages"] = n
        b = _Builder()
        with_text = failed = 0
        for i in range(min(n, MAX_PDF_PAGES)):
            try:
                t = reader.pages[i].extract_text() or ""
            except Exception as e:  # one broken page does not lose the others
                logger.debug("pdf %s page %d: %s", path.name, i + 1, e)
                t, failed = "", failed + 1
            t = t.replace("\r\n", "\n").strip()
            if t:
                with_text += 1
            b.segment("page", i + 1, f"--- 第 {i + 1} 頁 ---\n", t)
            if b.full:
                _truncated(info, f"只抽到第 {i + 1} 頁（{MAX_EXTRACT_CHARS:,} 字上限）")
                break
    if n > MAX_PDF_PAGES and not b.full:
        _truncated(info, f"共 {n} 頁，只抽了前 {MAX_PDF_PAGES} 頁")
    info["pages_with_text"] = with_text
    if failed:
        info["pages_failed"] = failed
    if n == 0:
        return unreadable("pdf", "empty", **info)
    if not with_text:
        return unreadable("pdf", "no_text", **info)
    return Extracted(type="pdf", readable=True, text=b.text, segments=b.segments, info=info)


def check_zip(path: Path) -> Optional[str]:
    """``None`` when a zip-based file is safe to open, else a reason
    (``too_large`` for a bomb, ``corrupt`` for a broken archive)."""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        return "corrupt"
    if len(infos) > MAX_ZIP_ENTRIES:
        return "too_large"
    total = 0
    for zi in infos:
        total += zi.file_size
        if total > MAX_ZIP_UNCOMPRESSED:
            return "too_large"
        if zi.file_size > 1024 * 1024 and zi.file_size > max(1, zi.compress_size) * MAX_ZIP_RATIO:
            return "too_large"
    return None


def _extract_docx(path: Path) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    b = _Builder()
    paragraphs = tables = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text
            if t.strip():
                paragraphs += 1
                b.add(t + "\n")
        elif tag == "tbl":
            tables += 1
            rows = []
            for row in Table(child, doc).rows:
                seen, cells = set(), []
                for cell in row.cells:  # a merged cell is listed once per grid column
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    cells.append(" ".join(cell.text.split()))
                rows.append(" | ".join(cells))
            b.add(f"[表格 {tables}]\n" + "\n".join(rows) + "\n\n")
        if b.full:
            break
    info: dict[str, Any] = {"paragraphs": paragraphs, "tables": tables}
    if b.full:
        _truncated(info, f"只取了前 {MAX_EXTRACT_CHARS:,} 字")
    if not b.text.strip():
        return unreadable("docx", "no_text", **info)
    return Extracted(type="docx", readable=True, text=b.text, info=info)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _extract_xlsx(path: Path) -> Extracted:
    from openpyxl import load_workbook

    # data_only: the values Excel saved, no formula is evaluated; keep_links=False: external links are not read
    wb = load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
    info: dict[str, Any] = {"sheets": []}
    b = _Builder()
    try:
        sheets = [ws for ws in wb.worksheets if hasattr(ws, "iter_rows")]
        if len(sheets) > MAX_SHEETS:
            _truncated(info, f"共 {len(sheets)} 個工作表，只讀了前 {MAX_SHEETS} 個")
        for idx, ws in enumerate(sheets[:MAX_SHEETS], 1):
            buf = io.StringIO()
            w = csv.writer(buf, lineterminator="\n")
            n, cut, cols = 0, False, 0
            pending_blank = 0
            for row in ws.iter_rows(max_col=MAX_SHEET_COLS, values_only=True):
                cells = [_cell(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                cols = max(cols, len(cells))
                if not cells:
                    pending_blank += 1  # trailing blank rows are dropped, inner ones kept
                    continue
                if n + pending_blank + 1 > MAX_SHEET_ROWS:
                    cut = True
                    break
                for _ in range(pending_blank):
                    w.writerow([])
                n += pending_blank + 1
                pending_blank = 0
                w.writerow(cells)
            sheet: dict[str, Any] = {"name": ws.title, "rows": n, "cols": cols}
            max_row = getattr(ws, "max_row", None)
            if cut:
                sheet["truncated"] = True
                if isinstance(max_row, int):
                    sheet["rows_total"] = max_row
                _truncated(info, f"工作表「{ws.title}」只取了前 {MAX_SHEET_ROWS:,} 列" + (f"（共約 {max_row:,} 列）" if isinstance(max_row, int) else ""))
            max_col = getattr(ws, "max_column", None)
            if isinstance(max_col, int) and max_col > MAX_SHEET_COLS:
                sheet["cols_truncated"] = True
                _truncated(info, f"工作表「{ws.title}」只取了前 {MAX_SHEET_COLS} 欄")
            info["sheets"].append(sheet)
            b.segment("sheet", idx, f"--- 工作表 {idx}：{ws.title}（{n} 列）---\n", buf.getvalue().rstrip("\n"), name=ws.title, rows=n)
            if b.full:
                _truncated(info, f"只取了前 {MAX_EXTRACT_CHARS:,} 字")
                break
    finally:
        wb.close()
    info["rows"] = sum(s["rows"] for s in info["sheets"])
    info["cols"] = info["sheets"][0]["cols"] if info["sheets"] else 0  # the first sheet's (the page shows one number)
    if not b.text.strip() or not info["rows"]:
        return unreadable("xlsx", "empty", **info)
    return Extracted(type="xlsx", readable=True, text=b.text, segments=b.segments, info=info)


def _shape_texts(shape: Any, out: list[str]) -> None:
    if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):  # a group
        for s in shape.shapes:
            _shape_texts(s, out)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        out.extend(" | ".join(" ".join(c.text.split()) for c in row.cells) for row in shape.table.rows)


def _extract_pptx(path: Path) -> Extracted:
    from pptx import Presentation

    prs = Presentation(str(path))
    b = _Builder()
    info: dict[str, Any] = {}
    slides = list(prs.slides)
    info["slides"] = len(slides)
    with_text = 0
    for i, slide in enumerate(slides[:MAX_SLIDES], 1):
        texts: list[str] = []
        for shape in slide.shapes:
            _shape_texts(shape, texts)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(texts) + (f"\n（備註）{notes}" if notes else "")
        if body.strip():
            with_text += 1
        b.segment("slide", i, f"--- 第 {i} 張投影片 ---\n", body.strip())
        if b.full:
            _truncated(info, f"只取到第 {i} 張（{MAX_EXTRACT_CHARS:,} 字上限）")
            break
    if len(slides) > MAX_SLIDES:
        _truncated(info, f"共 {len(slides)} 張，只讀了前 {MAX_SLIDES} 張")
    info["slides_with_text"] = with_text
    if not with_text:
        return unreadable("pptx", "no_text", **info)
    return Extracted(type="pptx", readable=True, text=b.text, segments=b.segments, info=info)


def _extract(path: Path, name: Optional[str]) -> Extracted:
    if path.stat().st_size == 0:
        return unreadable(detect(path, name), "empty")
    type_ = detect(path, name)
    if type_ in TEXT_TYPES:
        return _extract_text(path, type_)
    if type_ == "pdf":
        return _extract_pdf(path)
    if type_ in ("docx", "xlsx", "pptx"):
        bad = check_zip(path)
        if bad:
            return unreadable(type_, bad)
        try:
            return {"docx": _extract_docx, "xlsx": _extract_xlsx, "pptx": _extract_pptx}[type_](path)
        except Exception as e:  # a library that cannot parse it: a broken file
            logger.info("%s %s could not be read: %s", type_, path.name, e)
            return unreadable(type_, "corrupt")
    if type_ == "encrypted_office":
        return unreadable(Path(name or path.name).suffix.lower().lstrip(".")[:4] or "binary", "encrypted")
    if type_ == "legacy_office":
        return unreadable("legacy_office", "unsupported")
    if type_ == "audio":
        return unreadable("audio", "audio")
    if type_ == "image":
        return unreadable("image", "image")
    return unreadable("binary", "binary")


# ---------------------------------------------------------------- cache / entry points
_cache: OrderedDict[tuple[str, int, int], Extracted] = OrderedDict()
_lock = threading.Lock()


def extract_path(path: str | Path, name: Optional[str] = None) -> Extracted:
    """``path`` → its ``Extracted`` (cached by path, mtime, size). Blocking:
    call it from a worker thread (``aextract``). Never raises for what the
    file holds — an unreadable file is an ``Extracted`` with a reason; a
    missing file raises ``FileNotFoundError``."""
    p = Path(path)
    st = p.stat()
    key = (str(p), st.st_mtime_ns, st.st_size)
    with _lock:
        hit = _cache.get(key)
        if hit is not None:
            _cache.move_to_end(key)
            return hit
    try:
        out = _extract(p, name)
    except FileNotFoundError:
        raise
    except MemoryError:
        out = unreadable("binary", "too_large")
    except Exception as e:  # pragma: no cover - a reader bug must not break a chat turn
        logger.warning("extracting %s failed: %s", p.name, e)
        out = unreadable("binary", "corrupt")
    out.info["bytes"] = st.st_size
    with _lock:
        _cache[key] = out
        while len(_cache) > 1 and sum(len(v.text) for v in _cache.values()) > _CACHE_CHARS:
            _cache.popitem(last=False)
    return out


async def aextract(path: str | Path, name: Optional[str] = None, *, timeout: float = EXTRACT_TIMEOUT_S) -> Extracted:
    """``extract_path`` in a worker thread, giving up after ``timeout``
    seconds (the thread is left to finish and fill the cache)."""
    try:
        return await asyncio.wait_for(asyncio.to_thread(extract_path, path, name), timeout)
    except asyncio.TimeoutError:
        return unreadable("binary", "timeout")


def from_text(type_: str, text: Optional[str], **info: Any) -> Extracted:
    """A text work (lyrics, transcript) as an ``Extracted``."""
    text = (text or "").replace("\r\n", "\n")
    if len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS]
        _truncated(info, f"只取了前 {MAX_EXTRACT_CHARS:,} 字")
    if not text.strip():
        return unreadable(type_, "empty", **info)
    info.setdefault("lines", text.count("\n") + 1)
    return Extracted(type=type_, readable=True, text=text, info=info)


def public_info(ex: Extracted) -> dict[str, Any]:
    """``info`` as the API hands it out (plain JSON)."""
    out = {k: v for k, v in ex.info.items() if v is not None}
    if out.get("truncated_notes"):
        out["truncated_note"] = "；".join(out["truncated_notes"])
    return out


# ---------------------------------------------------------------- reading a part
def parse_range(raw: Any, top: int) -> Optional[tuple[int, int]]:
    """``"3-7"`` / ``"3"`` / ``3`` / ``[3, 7]`` → (3, 7) clipped to 1..top; ``None`` when unusable."""
    if raw is None or raw == "":
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        a = b = raw
    elif isinstance(raw, (list, tuple)) and raw and all(isinstance(x, int) for x in raw):
        a, b = raw[0], raw[-1]
    else:
        m = re.fullmatch(r"\s*(\d+)\s*(?:[-–~～到至,:]\s*(\d+)?\s*)?", str(raw))
        if not m:
            return None
        a = int(m.group(1))
        b = int(m.group(2)) if m.group(2) else (top if re.search(r"[-–~～到至,:]", str(raw)) else a)
    a, b = max(1, a), min(top, b)
    return (a, b) if a <= b else None


@dataclass
class Read:
    text: str
    label: str  # e.g. 第 3–7 頁 / 第 1–4,000 字 / 工作表「A」第 1–200 列
    start: int
    end: int
    more: Optional[str] = None  # how to ask for the next part, when one was cut


def _span_label(a: int, b: int, unit: str) -> str:
    return f"第 {a} {unit}" if a == b else f"第 {a}–{b} {unit}"


def _cap(ex: Extracted, start: int, end: int, cap: int) -> tuple[int, Optional[str]]:
    if end - start > cap:
        return start + cap, f"這一段超過 {cap:,} 字只給了前面；下一段用 start={start + cap}"
    return end, None


def read_range(ex: Extracted, *, start: Any = None, length: Any = None, lines: Any = None, pages: Any = None,
               sheet: Any = None, rows: Any = None, cap: int) -> Read:
    """One part of an extracted file, at most ``cap`` characters. Pages (PDF)
    / slides (pptx) by ``pages``; a sheet by ``sheet`` (name or 1-based
    number) and optionally ``rows`` within it; ``lines`` for anything; else
    ``start`` + ``length`` characters (default: from the beginning)."""
    text = ex.text
    if pages is not None and ex.segments and ex.segments[0].unit in ("page", "slide"):
        unit = "頁" if ex.segments[0].unit == "page" else "張投影片"
        top = ex.segments[-1].index
        span = parse_range(pages, top)
        if span is None:
            raise ValueError(f"pages 要是 1 到 {top} 之間的頁碼，例如 \"3\" 或 \"3-7\"")
        segs = [s for s in ex.segments if span[0] <= s.index <= span[1]]
        body = "\n\n".join((f"--- 第 {s.index} {'頁' if unit == '頁' else '張'} ---\n" + text[s.start:s.end]) for s in segs)
        more = None
        if len(body) > cap:
            body, more = body[:cap], f"這幾{unit}超過 {cap:,} 字只給了前面；請縮小範圍"
        return Read(body, _span_label(span[0], span[1], unit), segs[0].start, segs[-1].end, more)
    if (sheet is not None or rows is not None) and ex.segments and ex.segments[0].unit == "sheet":
        seg = None
        if sheet is None:
            seg = ex.segments[0]
        else:
            seg = next((s for s in ex.segments if s.name == str(sheet)), None)
            if seg is None and str(sheet).strip().isdigit():
                seg = next((s for s in ex.segments if s.index == int(str(sheet).strip())), None)
            if seg is None:
                seg = next((s for s in ex.segments if s.name.lower() == str(sheet).strip().lower()), None)
        if seg is None:
            raise ValueError("沒有這個工作表；可用的：" + "、".join(f"{s.index}.{s.name}" for s in ex.segments))
        body_all = text[seg.start:seg.end]
        row_starts = [0] + [m.end() for m in re.finditer("\n", body_all)]
        total = len(row_starts)
        span = parse_range(rows, total) if rows is not None else (1, total)
        if span is None:
            raise ValueError(f"rows 要是 1 到 {total} 之間，例如 \"1-200\"")
        a = row_starts[span[0] - 1]
        b = row_starts[span[1]] if span[1] < total else len(body_all)
        end, more = _cap(ex, seg.start + a, seg.start + b, cap)
        label = f"工作表「{seg.name}」" + (_span_label(span[0], span[1], "列") if rows is not None else "")
        return Read(text[seg.start + a:end].rstrip("\n"), label, seg.start + a, end, more)
    if lines is not None:
        starts = ex.line_starts()
        span = parse_range(lines, len(starts))
        if span is None:
            raise ValueError(f"lines 要是 1 到 {len(starts)} 之間，例如 \"100-200\"")
        a = starts[span[0] - 1]
        b = starts[span[1]] if span[1] < len(starts) else len(text)
        end, more = _cap(ex, a, b, cap)
        return Read(text[a:end].rstrip("\n"), _span_label(span[0], span[1], "行"), a, end, more)
    try:
        s = max(0, int(start or 0))
        n = int(length) if length not in (None, "") else cap
    except (TypeError, ValueError):
        raise ValueError("start 與 length 要是整數")
    if s >= len(text):
        raise ValueError(f"start 超過檔案長度（共 {len(text):,} 字）")
    n = max(1, min(n, cap))
    end = min(len(text), s + n)
    more = f"還有 {len(text) - end:,} 字；下一段用 start={end}" if end < len(text) else None
    return Read(text[s:end], f"第 {s + 1:,}–{end:,} 字", s, end, more)


def locate(ex: Extracted, offset: int) -> str:
    """Where ``offset`` is, in the words a person uses for this file."""
    for s in ex.segments:
        if s.start <= offset <= s.end:
            if s.unit == "page":
                return f"第 {s.index} 頁"
            if s.unit == "slide":
                return f"第 {s.index} 張投影片"
            row = ex.text.count("\n", s.start, offset) + 1
            return f"工作表「{s.name}」第 {row} 列"
    import bisect

    line = bisect.bisect_right(ex.line_starts(), offset)
    return f"第 {line} 行"


@dataclass
class Hit:
    where: str
    offset: int
    snippet: str


def search(ex: Extracted, query: str, *, max_hits: int, context: int = 80) -> tuple[list[Hit], int]:
    """Case-insensitive occurrences of ``query``: (the first ``max_hits``, how many in all)."""
    q = query.strip()
    if not q:
        return [], 0
    hay = ex.text.lower()
    needle = q.lower()
    hits: list[Hit] = []
    total, i = 0, hay.find(needle)
    while i != -1:
        total += 1
        if len(hits) < max_hits:
            a, b = max(0, i - context), min(len(ex.text), i + len(q) + context)
            snippet = ("…" if a else "") + ex.text[a:b].replace("\n", " ⏎ ") + ("…" if b < len(ex.text) else "")
            hits.append(Hit(locate(ex, i), i, snippet))
        i = hay.find(needle, i + max(1, len(needle)))
    return hits, total
