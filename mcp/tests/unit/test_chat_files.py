"""1.2-M5 檔案附件 (backend): extraction of every format (test files made
on the spot), encodings, broken / encrypted / oversized files, the delivery
decision per provider × file type × size, the shape of file parts in each
vendor's request (fake transports — no vendor is called), the in-turn file
tool loop, the transcription question before a reply, models without tools,
the upload endpoint (the generate page's behaviour unchanged) and the export."""

from __future__ import annotations

import asyncio
import io
import json
import time
import wave
import zipfile
from pathlib import Path
from types import SimpleNamespace as NS

import pytest

from omniapi_mcp.bus import EventBus
from omniapi_mcp.capabilities.text import AnthropicTextProvider, GeminiTextProvider, OpenAITextProvider
from omniapi_mcp.catalog import catalog
from omniapi_mcp.catalog.catalog import ModelEntry, discovered_input
from omniapi_mcp.chat import ChatError, ChatManager
from omniapi_mcp.chat import delivery as D
from omniapi_mcp.chat import files as F
from omniapi_mcp.chat import tools as T
from omniapi_mcp.chat.tools import ListFiles, ProposeGeneration, ReadFile, SearchFile, ToolRegistry, TranscriptionGate
from omniapi_mcp.generate import GenerationManager
from omniapi_mcp.providers.base import ProviderConfig
from omniapi_mcp.recorder import call_scope
from omniapi_mcp.store.db import Store
from omniapi_mcp.tools.text import TextTool


# =========================================================================== test files, made on the spot
def make_pdf(pages: list[str]) -> bytes:
    """A small valid PDF, one line of Helvetica text per page ("" = a page with no text layer)."""
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content = (f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else "0 0 1 rg 10 10 100 100 re f").encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for k, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{k} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def encrypt_pdf(raw: bytes, password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for p in PdfReader(io.BytesIO(raw)).pages:
        w.add_page(p)
    w.encrypt(user_password=password, owner_password=password + "-owner")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_paragraph("季度報告：營收成長 12%")
    d.add_paragraph("第二段說明")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "門市", "營收"
    t.cell(1, 0).text, t.cell(1, 1).text = "信義", "120"
    d.save(str(path))


def make_xlsx(path: Path, rows: int = 5) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "銷售"
    ws.append(["日期", "門市", "杯數"])
    for i in range(rows):
        ws.append([f"9/{i + 1}", "信義", i * 10])
    ws2 = wb.create_sheet("備註")
    ws2.append(["只有一列"])
    ws["E1"] = "=SUM(C2:C3)"  # a formula: never evaluated, its cached value (none here) is read
    wb.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "新品上市"
    s1.placeholders[1].text = "杯套設計三款"
    s1.notes_slide.notes_text_frame.text = "講者備註：先講價格"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "時程"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "週", "事項"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "W1", "打樣"
    prs.save(str(path))


def make_wav(seconds: float = 1.0, rate: int = 8000) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(b"\x00\x00" * int(seconds * rate))
    return buf.getvalue()


def write(tmp: Path, name: str, data: bytes | str) -> Path:
    p = tmp / name
    p.write_bytes(data.encode("utf-8") if isinstance(data, str) else data)
    return p


@pytest.fixture(autouse=True)
def _fresh_cache():
    F._cache.clear()
    yield
    F._cache.clear()


# =========================================================================== extraction
def test_text_files_and_their_encodings(tmp_path):
    ex = F.extract_path(write(tmp_path, "a.txt", "第一行\n第二行\n"))
    assert ex.readable and ex.type == "text" and ex.text == "第一行\n第二行\n" and ex.info["encoding"] == "utf-8" and ex.info["lines"] == 3
    bom = F.extract_path(write(tmp_path, "b.md", b"\xef\xbb\xbf# \xe6\xa8\x99\xe9\xa1\x8c"))
    assert bom.type == "markdown" and bom.text == "# 標題" and "BOM" in bom.info["encoding"]
    u16 = F.extract_path(write(tmp_path, "c.txt", "hello 世界".encode("utf-16")))
    assert u16.text == "hello 世界"
    u16le = F.extract_path(write(tmp_path, "d.log", "plain ascii log line\n".encode("utf-16-le")))  # no BOM: told by its NULs
    assert u16le.readable and u16le.text.startswith("plain ascii log")
    big5 = F.extract_path(write(tmp_path, "e.csv", "品名,數量\n珍珠奶茶,3\n".encode("cp950")))
    assert big5.readable and big5.info["encoding"] == "big5" and "珍珠奶茶" in big5.text and big5.info["rows"] == 2 and big5.info["cols"] == 2
    code = F.extract_path(write(tmp_path, "f.py", "def f():\n    return 1\n"))
    assert code.type == "code" and code.readable
    lossy = F.extract_path(write(tmp_path, "g.txt", ("正常的中文內容。" * 30).encode("utf-8") + b"\xff tail"))
    assert lossy.readable and lossy.info["lossy"] is True and lossy.info["encoding"] == "utf-8" and "正常的中文內容" in lossy.text
    assert not F.extract_path(write(tmp_path, "h.txt", b"")).readable


def test_a_huge_text_file_is_read_up_to_the_cap_and_says_so(tmp_path, monkeypatch):
    monkeypatch.setattr(F, "MAX_TEXT_READ_BYTES", 1000)
    ex = F.extract_path(write(tmp_path, "big.log", "x" * 5000))
    assert ex.readable and ex.chars == 1000 and ex.info["truncated"] and "只讀了前 1,000 bytes" in F.public_info(ex)["truncated_note"]
    monkeypatch.setattr(F, "MAX_TEXT_READ_BYTES", 10_000)
    monkeypatch.setattr(F, "MAX_EXTRACT_CHARS", 300)
    capped = F.extract_path(write(tmp_path, "big2.log", "y" * 5000))
    assert capped.chars == 300 and capped.info["truncated"]


def test_pdf_pages_text_scans_encryption_and_breakage(tmp_path):
    ex = F.extract_path(write(tmp_path, "r.pdf", make_pdf(["Revenue grew 12 percent", "", "Page three closing"])))
    assert ex.type == "pdf" and ex.readable and ex.info["pages"] == 3 and ex.info["pages_with_text"] == 2
    assert [s.index for s in ex.segments] == [1, 2, 3] and "Revenue grew" in ex.text[ex.segments[0].start:ex.segments[0].end]
    scan = F.extract_path(write(tmp_path, "scan.pdf", make_pdf(["", ""])))
    assert not scan.readable and scan.reason == "no_text" and scan.info["pages"] == 2 and "沒有可抽取的文字" in scan.info["reason_text"]
    locked = F.extract_path(write(tmp_path, "locked.pdf", encrypt_pdf(make_pdf(["secret text"]))))
    assert not locked.readable and locked.reason == "encrypted"
    broken = F.extract_path(write(tmp_path, "broken.pdf", b"%PDF-1.4\n garbage that is not a pdf"))
    assert not broken.readable and broken.reason in ("corrupt", "empty")
    renamed = F.extract_path(write(tmp_path, "report.txt", make_pdf(["Content wins over the name"])))
    assert renamed.type == "pdf" and renamed.readable  # sniffed, not trusted from the extension


def test_office_files(tmp_path):
    make_docx(tmp_path / "w.docx")
    doc = F.extract_path(tmp_path / "w.docx")
    assert doc.type == "docx" and doc.readable and "營收成長 12%" in doc.text and "門市 | 營收" in doc.text and doc.info["tables"] == 1
    make_xlsx(tmp_path / "s.xlsx")
    xl = F.extract_path(tmp_path / "s.xlsx")
    assert xl.type == "xlsx" and xl.readable and [s["name"] for s in xl.info["sheets"]] == ["銷售", "備註"]
    assert xl.info["sheets"][0]["rows"] == 6 and xl.info["sheets"][0]["cols"] == 3 and xl.info["cols"] == 3
    assert "9/3,信義,20" in xl.text and "SUM" not in xl.text  # data_only: a formula is never evaluated nor shown
    make_pptx(tmp_path / "p.pptx")
    pp = F.extract_path(tmp_path / "p.pptx")
    assert pp.type == "pptx" and pp.info["slides"] == 2 and "杯套設計三款" in pp.text and "講者備註" in pp.text and "W1 | 打樣" in pp.text


def test_big_spreadsheets_are_cut_and_say_where(tmp_path, monkeypatch):
    monkeypatch.setattr(F, "MAX_SHEET_ROWS", 3)
    make_xlsx(tmp_path / "big.xlsx", rows=10)
    xl = F.extract_path(tmp_path / "big.xlsx")
    first = xl.info["sheets"][0]
    assert first["rows"] == 3 and first["truncated"] and xl.info["truncated"] and "只取了前 3 列" in F.public_info(xl)["truncated_note"]


def test_zip_bombs_are_never_unpacked(tmp_path):
    p = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/worksheets/sheet1.xml", "0" * (8 * 1024 * 1024))  # 8 MiB of zeros compresses ~1000:1
    ex = F.extract_path(p)
    assert ex.type == "xlsx" and not ex.readable and ex.reason == "too_large"


def test_unreadable_kinds_have_reasons(tmp_path):
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    assert F.extract_path(write(tmp_path, "old.doc", ole)).reason == "unsupported"
    assert F.extract_path(write(tmp_path, "locked.docx", ole)).reason == "encrypted"  # an encrypted docx is an OLE box
    blob = F.extract_path(write(tmp_path, "model.skp", bytes(range(256)) * 20))
    assert blob.type == "binary" and blob.reason == "binary" and "二進位" in blob.info["reason_text"]
    assert F.extract_path(write(tmp_path, "not.docx", b"PK\x03\x04 broken zip")).reason in ("corrupt", "binary")


def test_results_are_cached_by_path_mtime_and_size(tmp_path, monkeypatch):
    p = write(tmp_path, "c.txt", "一")
    first = F.extract_path(p)
    calls = []
    monkeypatch.setattr(F, "_extract", lambda *a: calls.append(a) or F.unreadable("text", "empty"))
    assert F.extract_path(p) is first and calls == []
    p.write_text("二二", encoding="utf-8")  # changed file: extracted again
    F.extract_path(p)
    assert len(calls) == 1


async def test_extraction_gives_up_after_its_timeout(tmp_path, monkeypatch):
    def slow(*a):
        time.sleep(0.5)
        return F.unreadable("text", "empty")

    monkeypatch.setattr(F, "extract_path", slow)
    ex = await F.aextract(write(tmp_path, "s.txt", "x"), timeout=0.05)
    assert ex.reason == "timeout"


def test_audio_facts(tmp_path):
    info = F.audio_info(write(tmp_path, "a.wav", make_wav(2.5)))
    assert info["format"] == "wav" and info["duration_s"] == 2.5
    frame = bytes([0xFF, 0xFB, 0x90, 0x00]) + b"\x00" * 413  # MPEG-1 Layer III, 128 kbps
    mp3 = write(tmp_path, "a.mp3", frame * 100)
    info = F.audio_info(mp3)
    assert info["format"] == "mp3" and info["duration_estimated"] and abs(info["duration_s"] - mp3.stat().st_size * 8 / 128000) < 0.2
    assert F.audio_info(write(tmp_path, "a.m4a", b"\x00\x00\x00\x18ftypM4A ")).get("duration_s") is None


# =========================================================================== reading a part
def test_read_range_and_search(tmp_path):
    pdf = F.extract_path(write(tmp_path, "r.pdf", make_pdf(["alpha one", "beta two", "gamma three", "delta four"])))
    got = F.read_range(pdf, pages="2-3", cap=10_000)
    assert got.label == "第 2–3 頁" and "beta two" in got.text and "gamma three" in got.text and "alpha" not in got.text
    with pytest.raises(ValueError):
        F.read_range(pdf, pages="9", cap=100)
    make_xlsx(tmp_path / "s.xlsx", rows=8)
    xl = F.extract_path(tmp_path / "s.xlsx")
    rows = F.read_range(xl, sheet="銷售", rows="2-3", cap=10_000)
    assert rows.label == "工作表「銷售」第 2–3 列" and rows.text.splitlines()[0].startswith("9/1,") and len(rows.text.splitlines()) == 2
    assert F.read_range(xl, sheet="2", cap=10_000).text == "只有一列"
    txt = F.extract_path(write(tmp_path, "t.txt", "\n".join(f"line {i}" for i in range(1, 101))))
    assert F.read_range(txt, lines="10-11", cap=1000).text == "line 10\nline 11"
    part = F.read_range(txt, start=0, length=20, cap=1000)
    assert part.label == "第 1–20 字" and "下一段用 start=20" in part.more
    capped = F.read_range(txt, start=0, length=99999, cap=50)
    assert len(capped.text) == 50
    hits, total = F.search(pdf, "GAMMA", max_hits=5)
    assert total == 1 and hits[0].where == "第 3 頁" and "gamma three" in hits[0].snippet
    hits, total = F.search(xl, "信義", max_hits=3)
    assert total == 8 and len(hits) == 3 and hits[0].where.startswith("工作表「銷售」第 2 列")
    assert F.search(txt, "line 42", max_hits=5)[0][0].where == "第 42 行"


# =========================================================================== the decision
def _bf(tmp_path, name, data, kind="file", **kw) -> D.BranchFile:
    p = write(tmp_path, name, data)
    return D.BranchFile(key=f"m:{name}", id=f"id-{name}", ref={"upload_id": f"id-{name}"}, kind=kind, name=name, message_id="m",
                        path=p, bytes=p.stat().st_size, **kw)


async def _decide(f, support, budget=None, tools=True):
    await f.extract()
    return D.decide(f, support, budget or D.Budget(native_b64=support.budget), tools=tools)


ANTHROPIC = D.NativeSupport("anthropic", "document", False)
OPENAI = D.NativeSupport("openai", "file", False)
OPENROUTER_FILES = D.NativeSupport("openrouter", "file", True)
GOOGLE = D.NativeSupport("google", None, True)
DEEPSEEK = D.NativeSupport("deepseek", None, False)


@pytest.mark.parametrize("support,mode,fmt", [(ANTHROPIC, "native", "document"), (OPENAI, "native", "file"), (OPENROUTER_FILES, "native", "file"),
                                              (GOOGLE, "text", None), (DEEPSEEK, "text", None)])
async def test_a_pdf_per_provider(tmp_path, support, mode, fmt):
    d = await _decide(_bf(tmp_path, "r.pdf", make_pdf(["Quarterly revenue"])), support)
    assert d.mode == mode and d.meta.get("format") == fmt
    if mode == "native":
        part = d.parts[-1]
        assert part["type"] == "file" and part["file"]["file_data"].startswith("data:application/pdf;base64,") and part["file"]["filename"] == "r.pdf"
        assert "read_file" in d.parts[0]["text"]  # natively sent files stay readable through the tools
    else:
        assert "Quarterly revenue" in d.parts[0]["text"] and d.meta["chars"] > 0


async def test_native_pdfs_fall_back_to_text_past_the_page_and_byte_limits(tmp_path):
    f = _bf(tmp_path, "r.pdf", make_pdf(["a"] * 5))
    assert (await _decide(f, ANTHROPIC, D.Budget(native_b64=10**8, pdf_pages=4))).mode == "text"
    assert (await _decide(f, ANTHROPIC, D.Budget(native_b64=100))).mode == "text"
    scan = _bf(tmp_path, "scan.pdf", make_pdf(["", ""]))
    assert (await _decide(scan, ANTHROPIC)).mode == "native"  # a scan has no text, but a model that reads PDFs sees its pages
    d = await _decide(scan, GOOGLE)
    assert d.mode == "unreadable" and d.meta["reason"] == "no_text" and "不要假裝看過" in d.parts[0]["text"]
    locked = _bf(tmp_path, "locked.pdf", encrypt_pdf(make_pdf(["x"])))
    assert (await _decide(locked, ANTHROPIC)).mode == "unreadable"


async def test_text_excerpt_and_the_request_budget(tmp_path):
    small = _bf(tmp_path, "s.csv", "a,b\n1,2\n")
    d = await _decide(small, DEEPSEEK)
    assert d.mode == "text" and "全文如下" in d.parts[0]["text"] and "a,b\n1,2" in d.parts[0]["text"]
    big = _bf(tmp_path, "big.csv", "col1,col2\n" + "".join(f"{i},{'x' * 40}\n" for i in range(2000)))
    with_tools = await _decide(big, DEEPSEEK, tools=True)
    assert with_tools.mode == "excerpt" and "read_file" in with_tools.parts[0]["text"] and with_tools.meta["tools"] is True
    assert with_tools.meta["sent_chars"] <= D.EXCERPT_CHARS and with_tools.meta["rows"] == 2001 and 0 < with_tools.meta["sent_rows"] < 2001
    blind = await _decide(big, DEEPSEEK, tools=False)
    assert blind.mode == "excerpt" and "只看得到開頭" in blind.parts[0]["text"] and "說明只看了開頭" in blind.parts[0]["text"]
    assert "列" in D.describe_delivery(blind.meta)
    # the request's total: the newer file is inlined, the older drops to an excerpt
    a = _bf(tmp_path, "old.txt", "舊" * 20_000)
    b = _bf(tmp_path, "new.txt", "新" * 20_000)
    for f in (a, b):
        await f.extract()
    budget = D.Budget(chars=30_000, native_b64=0)
    assert D.decide(b, DEEPSEEK, budget, tools=True).mode == "text"
    assert D.decide(a, DEEPSEEK, budget, tools=True).mode == "excerpt"


async def test_audio_decisions(tmp_path):
    wav = _bf(tmp_path, "v.wav", make_wav(1), kind="audio")
    d = await _decide(wav, GOOGLE)
    assert d.mode == "native" and d.parts[-1]["type"] == "input_audio" and d.parts[-1]["input_audio"]["format"] == "wav"
    m4a = _bf(tmp_path, "v.m4a", b"\x00\x00\x00\x18ftypM4A " + b"\x00" * 100, kind="audio")
    unheard = await _decide(m4a, GOOGLE)
    assert unheard.mode == "unreadable" and "只收 wav、mp3" in unheard.meta["reason_text"] and "不要假裝聽過" in unheard.parts[0]["text"]
    assert (await _decide(_bf(tmp_path, "x.wav", make_wav(1), kind="audio"), ANTHROPIC)).meta["reason"] == "audio_unheard"
    told = _bf(tmp_path, "t.wav", make_wav(1), kind="audio", transcript_id="art-t", transcript_text="大家好，這是逐字稿。")
    t = await _decide(told, ANTHROPIC)
    assert t.mode == "text" and t.meta["transcribed"] and "這是逐字稿" in t.parts[0]["text"]
    assert D.describe_delivery(t.meta).startswith("轉成文字後送出（逐字稿")


async def test_missing_and_binary_files(tmp_path):
    gone = D.BranchFile(key="m:0", id="u0", ref={"upload_id": "u0"}, kind="file", name="gone.pdf", message_id="m", path=tmp_path / "nope.pdf")
    d = await _decide(gone, ANTHROPIC)
    assert d.mode == "unreadable" and d.meta["reason"] == "missing"
    blob = await _decide(_bf(tmp_path, "cup.skp", bytes(range(256)) * 10), ANTHROPIC)
    assert blob.mode == "unreadable" and "二進位" in blob.parts[0]["text"] and D.describe_delivery(blob.meta).startswith("讀不了")


def test_catalog_pdf_and_audio_flags():
    e = lambda p, caps: ModelEntry(id="m", provider=p, modality="text", capabilities=caps)  # noqa: E731
    assert e("anthropic", {"vision": True}).pdf and not e("anthropic", {"vision": True}).audio
    assert e("openai", {"vision": True}).pdf and not e("openai", {"vision": True}).audio
    assert not e("google", {"vision": True}).pdf and e("google", {"vision": True}).audio
    assert not e("deepseek", {"vision": True}).pdf and not e("deepseek", {}).audio
    assert e("openrouter", {"files": True, "audio_input": True}).pdf and not e("openrouter", {}).pdf
    caps = e("openai", {"vision": True}).to_dict()["capabilities"]
    assert caps["pdf"] is True and caps["audio"] is False
    assert discovered_input({"architecture": {"input_modalities": ["text", "file", "audio"]}}, "file") is True
    assert discovered_input({"architecture": {"input_modalities": ["text"]}}, "audio") is False
    assert discovered_input({}, "file") is None
    # the curated models: Gemini hears, Claude reads PDFs, DeepSeek neither
    assert catalog.get("gemini-3.8-flash", "google").audio and not catalog.get("gemini-3.8-flash", "google").pdf
    assert catalog.get("claude-sonnet-5", "anthropic").pdf


# =========================================================================== vendor request shapes (fake transports)
def _pdf_part(name="r.pdf"):
    return {"type": "file", "file": {"filename": name, "file_data": "data:application/pdf;base64,JVBERi0x"}}


def test_anthropic_gets_a_document_block():
    p = AnthropicTextProvider(ProviderConfig(api_key="k"))
    msgs = [{"role": "user", "content": [{"type": "text", "text": "看這份"}, _pdf_part(), {"type": "input_audio", "input_audio": {"data": "x", "format": "wav"}}]}]
    req = p._build_request("claude-sonnet-5", msgs, {})
    blocks = req["messages"][0]["content"]
    assert blocks[1] == {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": "JVBERi0x"}, "title": "r.pdf"}
    assert blocks[2]["type"] == "text" and "聽不到" in blocks[2]["text"]  # never sent audio: said as text


async def test_openai_and_gemini_get_the_parts_as_they_are():
    from openai.types.chat import ChatCompletionChunk

    class _S:
        def __init__(self, items):
            self.items = items

        def __aiter__(self):
            async def g():
                for i in self.items:
                    yield i
            return g()

    chunk = ChatCompletionChunk.model_validate({"id": "c", "object": "chat.completion.chunk", "created": 1, "model": "m",
                                               "choices": [{"index": 0, "delta": {"content": "ok"}, "finish_reason": "stop"}]})
    for cls, model, part in ((OpenAITextProvider, "gpt-5.4-mini", _pdf_part()),
                             (GeminiTextProvider, "gemini-3.8-flash", {"type": "input_audio", "input_audio": {"data": "UklGRg==", "format": "wav"}})):
        p = cls(ProviderConfig(api_key="k"))
        sent = {}

        async def create(**req):
            sent.update(req)
            return _S([chunk])

        p.client = NS(chat=NS(completions=NS(create=create)))
        msgs = [{"role": "user", "content": [{"type": "text", "text": "q"}, part]}]
        _ = [x async for x in p.stream(model, msgs)]
        assert sent["messages"][-1]["content"][1] == part


# =========================================================================== chat turns
def _settings_without_providers():
    off = NS(enabled=False, api_key="")
    return NS(providers=NS(openai=off, deepseek=off, anthropic=off, gemini=off, openrouter=off))


@pytest.fixture
def dev(monkeypatch):
    monkeypatch.setenv("OMNIAPI_DEV", "1")
    monkeypatch.setenv("OMNIAPI_ECHO_DELAY", "0")
    monkeypatch.delenv("OMNIAPI_OFFLINE", raising=False)


@pytest.fixture
async def store(tmp_path):
    s = Store(tmp_path / "chat.db")
    await s.open()
    yield s
    await s.close()


async def _no_menu(_ctx):
    return {"image": {"models": [{"id": "gpt-image-2", "name": "GPT Image 2", "pricing": None}], "default": "gpt-image-2", "sandbox": True},
            "speech": {"models": [], "default": None, "sandbox": True}}


async def _transcribe_default(_ctx):
    return "gpt-transcribe"


def _registry():
    reg = ToolRegistry()
    for t in (ProposeGeneration(menu=_no_menu), ListFiles(), ReadFile(), SearchFile(), TranscriptionGate(default=_transcribe_default)):
        reg.register(t)
    return reg


class FakeGenerations(GenerationManager):
    """The real job manager around a stand-in transcription tool."""

    def __init__(self, store, bus, tmp_path):
        super().__init__(store, bus)
        self.tmp = tmp_path
        self.gate: asyncio.Event | None = None
        self.fail: str | None = None
        self.n = 0

    def _validate(self, tool, args):
        return None

    def _tool(self, name):
        mgr = self

        async def run(args, convert_result=False):
            scope = call_scope.get()
            scope.call_id = await mgr.store.call_started(name, {}, source=scope.source, conversation_id=scope.links.get("conversation_id"))
            if mgr.gate is not None:
                await mgr.gate.wait()
            if mgr.fail:
                return {"error": mgr.fail}
            mgr.n += 1
            f = mgr.tmp / f"transcript{mgr.n}.txt"
            text = f"（示範逐字稿）第 {mgr.n} 份：大家好，今天討論新品杯套。"
            f.write_text(text, encoding="utf-8")
            row = await mgr.store.add_artifact(kind="transcript", tool=name, model=args.get("model"), file_path=str(f), text=text,
                                               source=scope.source, call_id=scope.call_id, meta=dict(scope.links))
            scope.artifacts.append(row)
            return {"model": args.get("model"), "cost_usd": 0.01}

        return NS(run=run)

    async def settle(self):
        await asyncio.gather(*list(self._tasks.values()), return_exceptions=True)


@pytest.fixture
async def env(store, dev, tmp_path):
    bus = EventBus()
    gens = FakeGenerations(store, bus, tmp_path)
    mgr = ChatManager(store, bus, TextTool(_settings_without_providers()))
    mgr.tools = _registry()
    await mgr.attach(gens)
    yield NS(chat=mgr, gens=gens, bus=bus, store=store, tmp=tmp_path)
    await mgr.close()
    await gens.close()


async def _upload(store, tmp, name, data, kind="file"):
    p = write(tmp, name, data)
    uid = f"up_{name.replace('.', '_')}"
    meta = {"info": F.audio_info(p)} if kind == "audio" else {"info": F.public_info(F.extract_path(p))} if kind == "file" else None
    await store.add_upload(upload_id=uid, kind=kind, filename=name, file_path=str(p), mime="application/octet-stream", size=p.stat().st_size, meta=meta)
    return uid


async def _settle(chat, cid):
    for _ in range(200):
        task = chat._tasks.get(cid)
        if task is None:
            return
        await asyncio.wait_for(task, 5)


async def test_files_reach_the_echo_model_and_the_reply_says_how(env):
    pdf = await _upload(env.store, env.tmp, "plan.pdf", make_pdf(["Launch plan for the new cup"]))
    csv = await _upload(env.store, env.tmp, "sales.csv", "date,cups\n9/1,120\n9/2,98\n")
    blob = await _upload(env.store, env.tmp, "cup.skp", bytes(range(256)) * 8)
    conv = await env.chat.create(model="echo-fast")
    res = await env.chat.send_and_wait(conv["id"], "這兩份在講什麼", attachments=[{"upload_id": pdf}, {"kind": "file", "upload_id": csv},
                                                                                   {"upload_id": blob}])
    msg = res["message"]
    assert res["state"] == "done" and "plan.pdf" in msg["content"] and "原樣附上" in msg["content"]
    assert "date,cups" in msg["content"] and "讀不了" in msg["content"] and "收到一個原樣的檔案 part：plan.pdf" in msg["content"]
    meta = msg["meta"]
    assert [(e["name"], e["mode"]) for e in meta["files"]] == [("plan.pdf", "native"), ("sales.csv", "text"), ("cup.skp", "unreadable")]
    assert meta["files_native"] == 1 and meta["files_text"] == 1 and meta["files_unreadable"] == 1 and meta["files_tools"] is True
    assert "images_sent" not in meta  # no images: no image report
    got = await env.chat.get(conv["id"])
    atts = got["messages"][0]["attachments"]
    assert [a["kind"] for a in atts] == ["file", "file", "file"] and atts[0]["info"]["pages"] == 1 and atts[1]["info"]["rows"] == 3
    assert atts[2]["info"]["readable"] is False and atts[2]["info"]["reason"] == "binary" and atts[0]["download_url"].endswith("download=true")
    assert got["title"] == "這兩份在講什麼"
    _, md = await env.chat.export_markdown(conv["id"])
    assert "> 附件 3 個：" in md and "plan.pdf：原樣送出" in md and "sales.csv：抽出文字" in md and "cup.skp：讀不了" in md


async def test_a_blind_model_without_tools_gets_text_and_excerpts(env):
    big = await _upload(env.store, env.tmp, "log.txt", "".join(f"第 {i} 行紀錄 {'x' * 30}\n" for i in range(3000)))
    pdf = await _upload(env.store, env.tmp, "plan.pdf", make_pdf(["Launch plan"]))
    conv = await env.chat.create(model="echo-blind")
    res = await env.chat.send_and_wait(conv["id"], "看一下", attachments=[{"upload_id": big}, {"upload_id": pdf}])
    meta = res["message"]["meta"]
    assert res["tools"] is False and meta["files_tools"] is False
    assert [(e["name"], e["mode"]) for e in meta["files"]] == [("log.txt", "excerpt"), ("plan.pdf", "text")]
    assert meta["files"][0]["tools"] is False and "只看得到開頭" in res["message"]["content"] and "Launch plan" in res["message"]["content"]


async def test_the_echo_model_reads_a_file_through_the_tool_loop(env):
    big = await _upload(env.store, env.tmp, "notes.txt", "開頭的一句話。" + "內容" * 20_000)
    conv = await env.chat.create(model="echo-fast")
    seen: list[dict] = []
    orig = env.bus.publish

    async def spy(ev, **kw):
        seen.append(ev)
        return await orig(ev, **kw)

    env.bus.publish = spy
    res = await env.chat.send_and_wait(conv["id"], "請讀這份檔案", attachments=[{"upload_id": big}])
    msg = res["message"]
    meta = msg["meta"]
    assert res["tools"] is True and meta["files"][0]["mode"] == "excerpt"
    assert meta["tool_rounds"] == 1 and meta["model_calls"] == 2 and meta["reads"][0]["tool"] == "read_file"
    assert meta["reads"][0]["name"] == "notes.txt" and meta["reads"][0]["what"] == "第 1–400 字" and meta["read_chars"] == 400
    assert msg["content"].startswith("我先用 read_file 讀一下檔案。\n\n收到第 **1** 輪") and "read_file 讀到：「notes.txt 第 1–400 字" in msg["content"]
    tools = [e for e in seen if e.get("type") == "chat.tool"]
    assert [e["state"] for e in tools] == ["running", "done"] and tools[0]["read"]["name"] == "notes.txt"
    deltas = "".join(e["delta"] for e in seen if e.get("type") == "chat.delta" and e["kind"] == "text")
    assert deltas == msg["content"]  # what streamed is what was stored
    calls = await env.store.calls(tool="chat")
    assert len(calls) == 1 and calls[0]["status"] == "ok"  # every round on one ledger row
    _, md = await env.chat.export_markdown(conv["id"])
    assert "> 讀過（1 輪工具" in md and "notes.txt 第 1–400 字" in md


class Looper:
    """A text tool scripted per round; routed as an echo model so it takes tools."""

    def __init__(self, rounds):
        self.rounds = rounds  # list of (text, calls)
        self.seen: list[list] = []
        self.gate: asyncio.Event | None = None

    def route(self, model):
        return model, NS(PROVIDER_KEY="echo")

    async def stream(self, messages, model=None, **params):
        self.seen.append((list(messages), params))
        text, calls = self.rounds[min(len(self.seen) - 1, len(self.rounds) - 1)]
        if self.gate is not None and len(self.seen) == 2:
            await self.gate.wait()
        yield {"type": "text", "delta": text}
        yield {"type": "done", "text": text, "model": model, "provider": "echo", "finish_reason": "tool_calls" if calls else "stop",
               "usage": {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15}, "cost_usd": 0.001, "tool_calls": calls}


def _read(cid, file_id, **args):
    return {"id": cid, "type": "function", "function": {"name": "read_file", "arguments": json.dumps({"file": file_id, **args})}}


async def _looper_chat(store, tmp, rounds):
    bus = EventBus()
    mgr = ChatManager(store, bus, Looper(rounds))
    mgr.tools = _registry()
    return mgr


async def test_the_loop_adds_up_usage_and_stops_at_the_round_limit(store, tmp_path, dev, monkeypatch):
    monkeypatch.setattr("omniapi_mcp.chat.manager.MAX_TOOL_ROUNDS", 3)
    pdf = await _upload(store, tmp_path, "r.pdf", make_pdf(["one", "two", "three"]))
    mgr = await _looper_chat(store, tmp_path, [("讀一下", [_read("c1", pdf, pages="2")])])  # it never stops asking
    conv = await mgr.create(model="echo-fast")
    res = await mgr.send_and_wait(conv["id"], "q", attachments=[{"upload_id": pdf}])
    meta = res["message"]["meta"]
    assert meta["tool_rounds"] == 3 and meta["model_calls"] == 4  # three rounds run, then one last call told to answer
    assert res["message"]["usage"] == {"prompt_tokens": 40, "completion_tokens": 20, "total_tokens": 60}
    assert res["message"]["cost_usd"] == pytest.approx(0.004)
    assert meta["tool_calls_ignored"] == [{"name": "read_file", "reason": "over the read limit of this turn"}]
    last_history = mgr.text.seen[-1][0]
    assert last_history[-1]["role"] == "tool" and T.LIMIT_NOTE in last_history[-1]["content"]
    assert [r["what"] for r in meta["reads"]] == ["第 2 頁"] * 3
    assert res["message"]["content"] == "\n\n".join(["讀一下"] * 4)
    assert res["message"].get("tool_calls") is None  # reads are never stored as calls
    await mgr.close()


async def test_a_proposal_in_a_reading_round_is_kept_for_the_end(store, tmp_path, dev):
    txt = await _upload(store, tmp_path, "t.txt", "內容")
    prop = {"id": "p1", "type": "function", "function": {"name": "propose_generation", "arguments": json.dumps({"kind": "image", "prompt": "杯套"})}}
    mgr = await _looper_chat(store, tmp_path, [("先讀", [_read("r1", txt), prop]), ("讀完了", None)])
    conv = await mgr.create(model="echo-fast")
    res = await mgr.send_and_wait(conv["id"], "讀完幫我畫", attachments=[{"upload_id": txt}])
    round2 = mgr.text.seen[1][0]
    results = {m["tool_call_id"]: m["content"] for m in round2 if m.get("role") == "tool"}
    assert "內容" in results["r1"] and results["p1"].startswith("（提議已記下")
    msg = res["message"]
    assert [c["function"]["name"] for c in msg["tool_calls"]] == ["propose_generation"]
    assert msg["proposals"][0]["state"] == "pending" and msg["meta"]["tool_rounds"] == 1
    await mgr.close()


async def test_a_turn_can_be_cancelled_in_the_middle_of_its_rounds(store, tmp_path, dev):
    txt = await _upload(store, tmp_path, "t.txt", "內容")
    mgr = await _looper_chat(store, tmp_path, [("先讀", [_read("r1", txt)]), ("讀完了", None)])
    mgr.text.gate = asyncio.Event()  # round two hangs
    conv = await mgr.create(model="echo-fast")
    await mgr.send(conv["id"], "q", attachments=[{"upload_id": txt}])
    for _ in range(100):
        if len(mgr.text.seen) >= 2:
            break
        await asyncio.sleep(0.01)
    assert (await mgr.cancel(conv["id"]))["cancelled"] is True
    reply = [m for m in await store.messages(conv["id"]) if m["role"] == "assistant"][-1]
    assert reply["meta"]["state"] == "cancelled" and reply["content"] == "先讀" and reply["meta"]["reads"][0]["tool"] == "read_file"
    assert reply["usage"]["total_tokens"] == 15  # the finished round is still counted
    assert not mgr.busy(conv["id"])
    await mgr.close()


async def test_tools_answer_bad_calls_without_breaking_the_turn(store, tmp_path, dev):
    txt = await _upload(store, tmp_path, "t.txt", "abc")
    bad = [_read("r1", "nope"), _read("r2", txt, pages="1"), {"id": "s1", "type": "function", "function": {"name": "search_file", "arguments": "{\"query\": \"b\"}"}},
           {"id": "l1", "type": "function", "function": {"name": "list_files", "arguments": "{}"}}, _read("r3", txt, start="x")]
    mgr = await _looper_chat(store, tmp_path, [("讀", bad), ("好", None)])
    conv = await mgr.create(model="echo-fast")
    res = await mgr.send_and_wait(conv["id"], "q", attachments=[{"upload_id": txt}])
    results = {m["tool_call_id"]: m["content"] for m in mgr.text.seen[1][0] if m.get("role") == "tool"}
    assert results["r1"].startswith("找不到這個檔案") and "t.txt（id:" in results["r1"]
    assert "第 1–3 字" in results["r2"]  # pages on a text file: read from the start
    assert "「b」共 1 處" in results["s1"] and "第 1 行" in results["s1"]
    assert "t.txt｜id:" in results["l1"] and "這次：全文已在訊息裡" in results["l1"]
    assert results["r3"].startswith("參數不對")
    assert res["state"] == "done" and len(res["message"]["meta"]["reads"]) == 5
    await mgr.close()


async def test_file_tools_are_offered_to_mcp_turns_with_files_but_proposals_are_not(store, tmp_path, dev):
    txt = await _upload(store, tmp_path, "t.txt", "abc")
    mgr = await _looper_chat(store, tmp_path, [("好", None)])
    conv = await mgr.create(model="echo-fast")
    await mgr.send_and_wait(conv["id"], "q", attachments=[{"upload_id": txt}], source="mcp")
    names = [t["function"]["name"] for t in mgr.text.seen[-1][1]["tools"]]
    assert names == ["list_files", "read_file", "search_file"]
    conv2 = await mgr.create(model="echo-fast")
    await mgr.send_and_wait(conv2["id"], "q", source="mcp")
    assert "tools" not in mgr.text.seen[-1][1]  # nothing to read, nobody to press a button
    await mgr.close()


async def test_works_as_attachments(env):
    lyr = await env.store.add_artifact(kind="lyrics", file_path=str(env.tmp / "l.txt"), text="[Verse]\n杯子裡的海")
    song = env.tmp / "song.wav"
    song.write_bytes(make_wav(1))
    music = await env.store.add_artifact(kind="music", file_path=str(song), title="夜曲", duration_s=1.0)
    conv = await env.chat.create(model="echo")
    res = await env.chat.send_and_wait(conv["id"], "歌詞跟歌", attachments=[{"artifact_id": lyr["id"]}, {"artifact_id": music["id"]}])
    meta = res["message"]["meta"]
    assert [(e["kind"], e["mode"]) for e in meta["files"]] == [("file", "text"), ("audio", "native")]
    assert "杯子裡的海" in res["message"]["content"] and "收到一段原樣的錄音 part，格式 wav" in res["message"]["content"]
    atts = (await env.chat.get(conv["id"]))["messages"][0]["attachments"]
    assert atts[0]["kind"] == "file" and atts[0]["info"]["chars"] > 0 and atts[1]["kind"] == "audio" and atts[1]["info"]["duration_s"] == 1.0
    with pytest.raises(ChatError):
        await env.chat.send(conv["id"], "x", attachments=[{"kind": "image", "artifact_id": lyr["id"]}])


# --------------------------------------------------------------------------- the transcription question
async def _audio_turn(env, model="echo-fast", name="meeting.m4a", data=None):
    up = await _upload(env.store, env.tmp, name, data or (b"\x00\x00\x00\x18ftypM4A " + b"\x00" * 64), kind="audio")
    conv = await env.chat.create(model=model)
    turn = await env.chat.send(conv["id"], "這段錄音在講什麼", attachments=[{"upload_id": up}])
    return conv["id"], up, turn


async def test_audio_the_model_cannot_hear_asks_first_then_replies_after_transcribing(env):
    cid, up, turn = await _audio_turn(env)
    assert turn["state"] == "awaiting" and not env.chat.busy(cid)
    p = turn["message"]["proposals"][0]
    assert p["kind"] == "transcript" and p["gate"] and p["state"] == "pending" and p["model"] == "gpt-transcribe"
    assert p["file"]["name"] == "meeting.m4a" and turn["message"]["content"] == "" and turn["message"]["meta"]["state"] == "awaiting"
    assert (await env.chat.wait(turn))["state"] == "awaiting"
    got = await env.chat.get(cid)
    assert got["messages"][-1]["meta"]["state"] == "awaiting" and got["messages"][-1]["proposals"][0]["state"] == "pending" and got["live"] is None
    assert await env.store.calls(tool="chat") == []  # no model was called

    accepted = await env.chat.accept(cid, p["id"])
    assert accepted["state"] == "generating"
    await env.gens.settle()
    await _settle(env.chat, cid)
    got = await env.chat.get(cid)
    reply = got["messages"][-1]
    assert reply["id"] == turn["message"]["id"] and reply["meta"]["state"] == "done" and reply["meta"]["asked_first"]
    assert reply["proposals"][0]["state"] == "done" and reply["proposals"][0]["artifact"]["kind"] == "transcript"
    assert reply["proposals"][0]["artifact"]["chars"] > 0
    assert reply["meta"]["files"][0]["mode"] == "text" and reply["meta"]["files"][0]["transcribed"] and "示範逐字稿" in reply["content"]
    assert (await env.store.upload(up))["meta"]["transcript_id"] == reply["proposals"][0]["artifact"]["id"]
    assert got["messages"][0]["attachments"][0]["transcript"]["chars"] > 0
    assert len([m for m in got["messages"] if m["role"] == "assistant"]) == 1  # filled in place: one reply
    # the next turn reads the transcript without asking again
    nxt = await env.chat.send_and_wait(cid, "再說一次")
    assert nxt["state"] == "done" and nxt["message"]["meta"]["files"][0]["mode"] == "text"
    _, md = await env.chat.export_markdown(cid)
    assert "回覆前先問：要不要把錄音 meeting.m4a 轉成文字" in md and "轉成文字後送出（逐字稿" in md


async def test_declining_replies_at_once_and_the_model_is_told_it_cannot_hear(env):
    cid, _, turn = await _audio_turn(env)
    p = turn["message"]["proposals"][0]
    await env.chat.decline(cid, p["id"])
    await _settle(env.chat, cid)
    reply = (await env.chat.get(cid))["messages"][-1]
    assert reply["meta"]["state"] == "done" and reply["proposals"][0]["state"] == "declined"
    f = reply["meta"]["files"][0]
    assert f["mode"] == "unreadable" and f["reason"] == "audio_unheard" and "聽不到" in reply["content"]
    with pytest.raises(ChatError) as e:
        await env.chat.decline(cid, p["id"])
    assert e.value.status == 409


async def test_a_failed_transcription_can_be_retried_or_skipped(env):
    cid, _, turn = await _audio_turn(env)
    p = turn["message"]["proposals"][0]
    env.gens.fail = "quota exceeded"
    await env.chat.accept(cid, p["id"])
    await env.gens.settle()
    got = (await env.chat.get(cid))["messages"][-1]
    assert got["proposals"][0]["state"] == "failed" and got["meta"]["state"] == "awaiting" and env.chat._tasks.get(cid) is None
    env.gens.fail = None
    await env.chat.decline(cid, p["id"])  # 不轉錄，直接回覆 after a failure
    await _settle(env.chat, cid)
    assert (await env.chat.get(cid))["messages"][-1]["meta"]["state"] == "done"


async def test_two_audio_files_two_questions_and_the_reply_waits_for_both(env):
    a = await _upload(env.store, env.tmp, "a.m4a", b"\x00\x00\x00\x18ftypM4A ", kind="audio")
    b = await _upload(env.store, env.tmp, "b.m4a", b"\x00\x00\x00\x18ftypM4A x", kind="audio")
    conv = await env.chat.create(model="echo-fast")
    turn = await env.chat.send(conv["id"], "兩段", attachments=[{"upload_id": a}, {"upload_id": b}])
    ps = turn["message"]["proposals"]
    assert len(ps) == 2
    await env.chat.decline(conv["id"], ps[0]["id"])
    assert env.chat._tasks.get(conv["id"]) is None and (await env.chat.get(conv["id"]))["messages"][-1]["meta"]["state"] == "awaiting"
    await env.chat.accept(conv["id"], ps[1]["id"])
    await env.gens.settle()
    await _settle(env.chat, conv["id"])
    reply = (await env.chat.get(conv["id"]))["messages"][-1]
    assert reply["meta"]["state"] == "done" and [e["mode"] for e in reply["meta"]["files"]] == ["unreadable", "text"]


async def test_moving_on_skips_the_waiting_reply(env):
    cid, _, turn = await _audio_turn(env)
    nxt = await env.chat.send_and_wait(cid, "算了，換個問題")
    assert nxt["state"] == "done"
    rows = {m["id"]: m for m in await env.store.messages(cid)}
    waited = rows[turn["message"]["id"]]
    assert waited["meta"]["state"] == "skipped" and waited["meta"]["gate"][turn["message"]["proposals"][0]["id"]]["auto"] is True
    with pytest.raises(ChatError) as e:
        await env.chat.accept(cid, turn["message"]["proposals"][0]["id"])
    assert e.value.status == 409
    # regenerate / edit still work around it; the history skips the empty reply
    again = await env.chat.wait(await env.chat.regenerate(cid, nxt["message"]["id"]))
    assert again["state"] == "done"


async def test_a_model_that_hears_gets_the_audio_without_asking(env):
    cid, _, turn = await _audio_turn(env, model="echo", name="v.wav", data=make_wav(1))
    assert turn["state"] == "streaming"
    await _settle(env.chat, cid)
    reply = (await env.chat.get(cid))["messages"][-1]
    assert reply["meta"]["files"][0]["mode"] == "native" and reply["meta"]["files"][0]["format"] == "input_audio"


async def test_a_waiting_reply_survives_a_restart(env, tmp_path):
    cid, _, turn = await _audio_turn(env)
    p = turn["message"]["proposals"][0]
    # simulate the process stopping right after the owner declined, before the reply ran
    reply = await env.store.message(turn["message"]["id"])
    meta = reply["meta"]
    meta["gate"][p["id"]]["state"] = "declined"
    await env.store.update_message(reply["id"], meta=meta)
    mgr2 = ChatManager(env.store, EventBus(), TextTool(_settings_without_providers()))
    mgr2.tools = _registry()
    await mgr2.attach(env.gens)  # reconcile resumes it
    await _settle(mgr2, cid)
    assert (await mgr2.get(cid))["messages"][-1]["meta"]["state"] == "done"
    await mgr2.close()


async def test_a_transcription_interrupted_by_a_restart_becomes_failed(env):
    cid, _, turn = await _audio_turn(env)
    p = turn["message"]["proposals"][0]
    env.gens.gate = asyncio.Event()
    await env.chat.accept(cid, p["id"])
    gid = (await env.chat.get(cid))["messages"][-1]["proposals"][0]["generation_id"]
    for t in list(env.gens._tasks.values()):
        t.cancel()
    await asyncio.gather(*list(env.gens._tasks.values()), return_exceptions=True)
    # the record is still "generating" if the callback raced; force the restart path
    await env.store.update_generation(gid, status="interrupted")
    reply = await env.store.message(turn["message"]["id"])
    m = reply["meta"]
    m["gate"][p["id"]]["state"] = "generating"
    await env.store.update_message(reply["id"], meta=m)
    await env.chat.reconcile()
    after = (await env.chat.get(cid))["messages"][-1]
    assert after["proposals"][0]["state"] == "failed" and after["meta"]["state"] == "awaiting"


# =========================================================================== the upload endpoint
@pytest.fixture
async def ctx(tmp_path):
    store = Store(tmp_path / "a.db")
    await store.open()
    storage = NS(base_path=str(tmp_path / "storage"))
    yield NS(store=store, bus=EventBus(), settings=NS(storage=storage))
    await store.close()


@pytest.fixture
def client(ctx, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from omniapi_mcp.daemon import app as daemon_app
    from omniapi_mcp.runtime import runtime

    monkeypatch.setenv("OMNIAPI_HOME", str(tmp_path / "home"))
    monkeypatch.setattr(runtime, "context", ctx)
    return TestClient(daemon_app.create_app(ctx.settings), base_url="http://127.0.0.1")  # loopback Host (1.3-M2 guard)


async def test_the_generate_page_still_gets_images_and_audio_only(ctx, client):
    def go():
        assert client.post("/api/uploads", params={"filename": "evil.exe"}, content=b"MZ").status_code == 415
        assert client.post("/api/uploads", params={"filename": "r.pdf"}, content=make_pdf(["x"])).status_code == 415
        assert client.post("/api/uploads", params={"filename": "fake.png"}, content=b"not an image").status_code == 400
        up = client.post("/api/uploads", params={"filename": "a.wav"}, content=make_wav(1)).json()
        assert up["kind"] == "audio" and up["info"] is None  # the generate page's uploads are not extracted
        assert client.post("/api/uploads", params={"filename": "x.txt", "purpose": "other"}, content=b"x").status_code == 422

    await asyncio.to_thread(go)


async def test_chat_uploads_take_any_file_and_say_what_it_holds(ctx, client, monkeypatch):
    from omniapi_mcp.daemon import app as daemon_app

    def go():
        pdf = client.post("/api/uploads", params={"filename": "plan.pdf", "purpose": "chat"}, content=make_pdf(["a", "b"])).json()
        assert pdf["kind"] == "file" and pdf["mime"] == "application/pdf" and pdf["info"]["pages"] == 2 and pdf["info"]["readable"] is True
        assert pdf["info"]["type_name"] == "PDF" and "file_path" not in pdf
        exe = client.post("/api/uploads", params={"filename": "tool.exe", "purpose": "chat"}, content=b"MZ\x90\x00" + bytes(range(256)) * 4).json()
        assert exe["kind"] == "file" and exe["info"]["readable"] is False and exe["info"]["reason"] == "binary"
        csv = client.post("/api/uploads", params={"filename": "s.csv", "purpose": "chat"}, content="a,b,c\n1,2,3\n".encode()).json()
        assert csv["info"]["rows"] == 2 and csv["info"]["cols"] == 3
        wav = client.post("/api/uploads", params={"filename": "v.wav", "purpose": "chat"}, content=make_wav(2)).json()
        assert wav["kind"] == "audio" and wav["info"]["duration_s"] == 2.0
        bad = client.post("/api/uploads", params={"filename": "fake.png", "purpose": "chat"}, content=b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 2).json()
        assert bad["kind"] == "file" and bad["info"]["readable"] is False  # taken, and said to be unreadable
        page = client.post("/api/uploads", params={"filename": "x.html", "purpose": "chat"}, content=b"<script>alert(1)</script>").json()
        r = client.get(page["file_url"])
        assert r.headers["content-type"] == "application/octet-stream" and "attachment" in r.headers["content-disposition"]
        assert r.headers["x-content-type-options"] == "nosniff"  # an uploaded page never renders on this origin
        assert client.get(f"/api/uploads/{pdf['id']}").json()["info"]["pages"] == 2
        limits = client.get("/api/uploads/limits").json()
        assert limits["max_bytes"]["file"] == 50 * 1024 * 1024 and limits["max_attachments"] == 10
        monkeypatch.setattr(daemon_app, "UPLOAD_MAX_FILE", 10)
        assert client.post("/api/uploads", params={"filename": "big.bin", "purpose": "chat"}, content=b"x" * 11).status_code == 413
        weird = client.post("/api/uploads", params={"filename": "a.péf", "purpose": "chat"}, content=b"x").json()
        assert weird["filename"] == "a.péf"  # the name is kept for display; the disk name has no odd extension

    await asyncio.to_thread(go)
    stored = [p.name for p in (Path(ctx.settings.storage.base_path) / "uploads").rglob("*") if p.is_file()]
    assert all(n.startswith("upload_") for n in stored) and not any(n.endswith("éf") for n in stored)


async def test_a_slow_extraction_answers_later(ctx, monkeypatch, tmp_path):
    import httpx

    from omniapi_mcp.daemon import app as daemon_app
    from omniapi_mcp.runtime import runtime

    monkeypatch.setenv("OMNIAPI_HOME", str(tmp_path / "home"))
    monkeypatch.setattr(runtime, "context", ctx)
    monkeypatch.setattr(daemon_app, "UPLOAD_EXTRACT_WAIT_S", 0.01)
    real = F.extract_path

    def slow(*a):
        time.sleep(0.3)
        return real(*a)

    monkeypatch.setattr(F, "extract_path", slow)
    app = daemon_app.create_app(ctx.settings)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://127.0.0.1") as http:  # this loop: background work goes on
        up = (await http.post("/api/uploads", params={"filename": "t.txt", "purpose": "chat"}, content="一二三".encode())).json()
        assert up["info"] is None and up["info_pending"] is True
        for _ in range(100):
            row = (await http.get(f"/api/uploads/{up['id']}")).json()
            if row.get("info"):
                break
            await asyncio.sleep(0.05)
        assert row["info"]["chars"] == 3 and "info_pending" not in row
